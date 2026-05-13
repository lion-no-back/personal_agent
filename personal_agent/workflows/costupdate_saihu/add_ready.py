import pandas as pd
import glob
import os
import math
import openpyxl
import logging

class ConsoleFilter(logging.Filter):
    def filter(self, record):
        if record.name == 'py.warnings':
            return False
        return True

def setup_logger(log_file):
    logging.captureWarnings(True)
    logger = logging.getLogger('add_ready_logger')
    logger.setLevel(logging.DEBUG)
    
    if logger.hasHandlers():
        logger.handlers.clear()

    console_handler = logging.StreamHandler()
    console_handler.setLevel(logging.INFO)
    console_handler.addFilter(ConsoleFilter())
    console_format = logging.Formatter('%(message)s')
    console_handler.setFormatter(console_format)

    file_handler = logging.FileHandler(log_file, encoding='utf-8', mode='w')
    file_handler.setLevel(logging.DEBUG)
    file_format = logging.Formatter('%(asctime)s | %(levelname)-8s | %(message)s', datefmt='%Y-%m-%d %H:%M:%S')
    file_handler.setFormatter(file_format)

    logger.addHandler(console_handler)
    logger.addHandler(file_handler)
    return logger

def process_and_split_data(template_path, name_csv_path, error_files_pattern, output_dir, log_file):
    logger = setup_logger(log_file)
    logger.info("========== 🚀 任务开始：赛狐报错记录清洗与重新生成 ==========")
    
    generated_files = [] # 用于收集生成的文件路径

    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    error_files = glob.glob(error_files_pattern)
    if not error_files:
        logger.error("❌ 未找到任何报错记录文件。")
        return generated_files

    error_df_list = []
    for file in error_files:
        try:
            df = pd.read_excel(file)
            error_df_list.append(df)
        except Exception as e:
            logger.error(f"  ├─ 读取失败: {file}，错误: {e}")
    
    merged_error_df = pd.concat(error_df_list, ignore_index=True)
    error_sku_col = next((col for col in merged_error_df.columns if 'SKU' in col.upper()), '*SKU')
    error_cost_col = next((col for col in merged_error_df.columns if '采购成本' in col), None)

    if not error_cost_col:
        logger.error("❌ 严重错误：未找到包含“采购成本”的列。")
        return generated_files

    merged_error_df[error_sku_col] = merged_error_df[error_sku_col].astype(str).str.strip()
    merged_error_df[error_sku_col] = merged_error_df[error_sku_col].apply(lambda x: x[:-2] if x.endswith('.0') else x)

    try:
        name_df = pd.read_csv(name_csv_path, dtype=str)
        name_df.columns = name_df.columns.str.strip()
        name_df['MSKU'] = name_df['MSKU'].str.strip()
        name_mapping = dict(zip(name_df['MSKU'], name_df['本地SKU名称']))
    except Exception as e:
        logger.error(f"❌ 读取映射表失败。错误: {e}")
        return generated_files

    master_data = []
    for _, row in merged_error_df.iterrows():
        sku_val = row[error_sku_col]
        name_val = name_mapping.get(sku_val, "")
        master_data.append({
            'sku': sku_val, 'cost': row[error_cost_col], 
            'name': name_val, 'status': "在售" if name_val else ""
        })

    chunk_size = 5100
    total_files = math.ceil(len(master_data) / chunk_size)

    for i in range(total_files):
        chunk_data = master_data[i * chunk_size : (i + 1) * chunk_size]
        output_filename = os.path.join(output_dir, f'批量导入_分割结果_第{i+1}部分.xlsx')
        
        try:
            wb = openpyxl.load_workbook(template_path)
            sheet = wb.active
            
            header_row = sheet[1]
            col_idx_sku = col_idx_name = col_idx_status = col_idx_cost = None
            
            for cell in header_row:
                if cell.value:
                    val_str = str(cell.value).strip()
                    if 'SKU' in val_str.upper(): col_idx_sku = cell.column
                    elif '品名' in val_str: col_idx_name = cell.column
                    elif '状态' in val_str: col_idx_status = cell.column
                    elif '采购成本' in val_str: col_idx_cost = cell.column

            for row_offset, data_item in enumerate(chunk_data):
                current_row = row_offset + 2
                if col_idx_sku: sheet.cell(row=current_row, column=col_idx_sku, value=data_item['sku'])
                if col_idx_name: sheet.cell(row=current_row, column=col_idx_name, value=data_item['name'])
                if col_idx_status: sheet.cell(row=current_row, column=col_idx_status, value=data_item['status'])
                if col_idx_cost: sheet.cell(row=current_row, column=col_idx_cost, value=data_item['cost'])

            wb.save(output_filename)
            wb.close()
            generated_files.append(output_filename) # 成功后加入列表
            
        except Exception as e:
            logger.error(f"  ├─ [X] 保存 第{i+1}部分 失败！错误信息: {e}")

    logger.info(f"✔ 成功导出 {len(generated_files)} 个文件。")
    return generated_files # 返回生成的文件列表给主控