import os
import sys
import asyncio
import aiohttp
import requests
import json
import logging
import re
import html
import time
import warnings
import pandas as pd
from urllib.parse import unquote
from tqdm.asyncio import tqdm

# ==========================================
# 🚨 强制依赖与警告控制
# ==========================================
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

try:
    import lxml
    from bs4 import BeautifulSoup
    import openpyxl 
except ImportError:
    print("\n❌ 严重错误：未安装必备的 lxml, bs4 或 openpyxl！")
    sys.exit(1)

# ==========================================
# 📝 1. 日志与全局配置
# ==========================================
class FlushFileHandler(logging.FileHandler):
    def emit(self, record):
        super().emit(record)
        self.flush()

logger = logging.getLogger("MaBangSpider")
logger.setLevel(logging.INFO)
logger.propagate = False

BASE_URL = "https://private.mabangerp.com"
AMZ_BASE_URL = "https://private-amz.mabangerp.com"

CHECK_TYPE_API = f"{AMZ_BASE_URL}/index.php?mod=fbanew.listsearch"
STOCK_API_URL = f"{AMZ_BASE_URL}/index.php?mod=stock.getStockList"
EXPORT_URL = f"{BASE_URL}/index.php?mod=combosku.doExportFileNew"

file_lock = asyncio.Lock()

CONCURRENCY_LIMIT = 40           
MAX_RETRIES = 3                  
BATCH_SIZE_TYPE = 300            
BATCH_SIZE_STOCK = 200
TIMEOUT_SETTING = aiohttp.ClientTimeout(total=120, sock_connect=10, sock_read=90)
WRITE_FLUSH_EVERY = 50           


# ==========================================
# 🔑 2. 鉴权中心
# ==========================================
def get_global_auth():
    print("  ⏳ 正在获取马帮跨域全局授权...")
    logger.info("=== 启动程序，开始获取鉴权 ===")
    
    main_session = requests.Session()
    base_headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        "Connection": "keep-alive"
    }
    
    try: main_session.get(f"{BASE_URL}/", headers=base_headers, timeout=10)
    except Exception: pass
        
    login_url = f"{BASE_URL}/index.php?mod=main.doLogin&lang=undefined"
    login_headers = base_headers.copy()
    login_headers.update({
        "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
        "X-Requested-With": "XMLHttpRequest",
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "Origin": BASE_URL,
    })
    
    data = {"username": os.environ.get("MABANG_USERNAME", ""), "password": os.environ.get("MABANG_PASSWORD", ""), "verifyType": "1"}
    resp = main_session.post(login_url, headers=login_headers, data=data, timeout=10)
    if not resp.json().get("success"):
        logger.error("主站登录失败，请检查密码")
        return None, None, None

    for label, url, method, payload in [
        ("GET /", f"{BASE_URL}/", "GET", None),
        ("POST main.index", f"{BASE_URL}/index.php?mod=main.index", "POST", {}),
        ("POST main.getUserInfo", f"{BASE_URL}/index.php?mod=main.getUserInfo", "POST", {}),
    ]:
        if method == "GET": main_session.get(url, headers=base_headers, timeout=10)
        else: main_session.post(url, headers=login_headers, data=payload, timeout=10)
        if any(c.name == "memberInfo" for c in main_session.cookies): break

    cookies_dict = requests.utils.dict_from_cookiejar(main_session.cookies)
    memcache_key = cookies_dict.get("MABANG_ERP_PRO_MEMBERINFO_LOGIN_COOKIE")
    if not memcache_key:
        logger.error("未找到 memcache_key，登录可能未完全成功")
        return None, None, None

    main_session.cookies.set("exportv2", "2", domain="private.mabangerp.com")
    main_session.cookies.set("login_js_cookie_phone_cookie", "", domain="private.mabangerp.com")

    amz_session = requests.Session()
    for c in main_session.cookies:
        amz_session.cookies.set(c.name, c.value, domain="private-amz.mabangerp.com")
        
    transition_url = f"{AMZ_BASE_URL}/index.php?mod=fbanew.list&platform=amazon&version=1&cMKey={memcache_key}&fromTest=111&lang=cn&tz=UTC+8&DOMAIN=private.mabangerp.com"
    try:
        amz_session.get(transition_url, headers=base_headers, timeout=15)
        amz_post_url = f"{AMZ_BASE_URL}/index.php?mod=combosku.getCombosSkuList"
        amz_post_headers = {
            **base_headers,
            "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
            "Origin": AMZ_BASE_URL,
            "X-Requested-With": "XMLHttpRequest",
        }
        amz_session.post(amz_post_url, headers=amz_post_headers, data="searchLike=comboSku&operate=Like", timeout=15)
        
        final_amz_cookies = requests.utils.dict_from_cookiejar(amz_session.cookies)
        amz_cookie_str = "; ".join([f"{k}={v}" for k, v in final_amz_cookies.items()])
        print("  ✅ 全局跨域授权成功，引擎已就绪！\n")
        return main_session, amz_cookie_str, memcache_key
    except Exception as e:
        logger.error(f"激活子站失败: {e}")
        return None, None, None


# ==========================================
# ⚡ 3. 核心流水线各阶段业务逻辑
# ==========================================
def _parse_type_html_fast(html_content, msku_batch):
    html_content = html.unescape(html_content)
    result = {m: "判断超时" for m in msku_batch}
    found = set()
    target_skus = set(msku_batch)
    
    rows = html_content.split('<tr ')
    for row in rows[1:]:
        msku_match = re.search(r'data-msku="([^"]+)"', row)
        if not msku_match: msku_match = re.search(r'msku="([^"]+)"', row)
            
        if msku_match:
            row_msku = msku_match.group(1).strip()
            if row_msku in target_skus:
                is_combo = ('class="label label-success grouplabel"' in row) or ('>组<' in row) or ('grouplabel' in row and '组' in row)
                if row_msku not in found:
                    result[row_msku] = "组合SKU" if is_combo else "普通SKU"
                    found.add(row_msku)
                elif is_combo and result[row_msku] == "普通SKU":
                    result[row_msku] = "组合SKU"
                
    for m in msku_batch:
        if m not in found and result[m] == "判断超时":
            result[m] = "无此SKU"
    return result

async def batch_check_sku_types(session, msku_batch):
    platform_sku_data = "\r\n".join(msku_batch)
    form_data = aiohttp.FormData()
    form_data.add_field('shopId', '')
    form_data.add_field('status', '1')
    form_data.add_field('isChange', '1')
    form_data.add_field('searchtexttype', 'platformSkuLike')
    form_data.add_field('atn', 'list')
    form_data.add_field('searchtype', '1')
    form_data.add_field('page', '1')
    form_data.add_field('rowsPerPage', str(max(len(msku_batch) * 2, 300)))
    form_data.add_field('selecttype', 'platformSkuIn')
    form_data.add_field('platformSkuData', platform_sku_data)

    result = {m: "判断超时" for m in msku_batch}
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            async with session.post(CHECK_TYPE_API, data=form_data, timeout=TIMEOUT_SETTING) as resp:
                if resp.status != 200: continue
                data = await resp.json(content_type=None)
                if not data.get('success'): continue
                html_content = data.get('html', '')
                if not html_content.strip() or '暂无数据' in html_content:
                    return {m: "无此SKU" for m in msku_batch}
                return await asyncio.to_thread(_parse_type_html_fast, html_content, msku_batch)
        except Exception:
            if attempt < MAX_RETRIES: await asyncio.sleep(2 * attempt)
    return result


# 【阶段 C1】增强版：防御网络崩溃 + 分批导出机制 + 超时放宽
def sync_export_combo_and_parse(main_session, memcache_key, combo_local_skus_unique, output_dir):
    if not combo_local_skus_unique: return {}
    
    export_headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
        "X-Requested-With": "XMLHttpRequest",
        "Origin": BASE_URL,
    }
    
    # 核心修复 1：将庞大的 SKU 列表按 3000 一批进行切割
    BATCH_SIZE = 3000
    sku_batches = [combo_local_skus_unique[i:i + BATCH_SIZE] for i in range(0, len(combo_local_skus_unique), BATCH_SIZE)]
    
    final_combo_dict = {}

    for batch_idx, current_batch in enumerate(sku_batches):
        print(f"\n     ↳ 正在处理组合导出批次 {batch_idx + 1}/{len(sku_batches)} (本批 {len(current_batch)} 条)...")
        
        step1_data = [
            ("backUrl", ""),
            ("orderIds", "\r\n".join(current_batch)),
            ("fieldlabel", "uq101"),
            ("fieldlabel", "uq136"),
            ("fieldlabel", "uq138"),
            ("map-name[]", "组合sku编码"),
            ("map-uq[]", "uq101"),
            ("map-text[]", ""),
            ("map-name[]", "关联sku个数"),
            ("map-uq[]", "uq136"),
            ("map-text[]", ""),
            ("map-name[]", "关联sku信息"),
            ("map-uq[]", "uq138"),
            ("map-text[]", ""),
            ("templateName", ""),
            ("templateId", "0"),
            ("datasOpen", "1"),
            ("memcacheKey", memcache_key),
            ("showRmbColumn", "2"),
            ("pageSave", "1"),
            ("operateType", "19"),
            ("params", ""),
            ("InterfaceUrl", ""),
            ("mainMenu", ""),
            ("hiddenPage", "1"),
            ("hiddenPageSize", ""),
            ("tableBase", ""),
            ("isMerage", "1"),
            ("version", "v2"),
            ("step", "1"),
        ]
        
        try:
            # 核心修复 2：将所有相关的 timeout 统一放宽到 150 秒
            r1 = main_session.post(EXPORT_URL, headers=export_headers, data=step1_data, timeout=150)
            j1 = r1.json()
            if not j1.get("success"): 
                print(f"     ❌ [警告] 批次 {batch_idx + 1} 首步被拒: {j1.get('message', j1)}")
                continue # 失败则跳过本批次，不让程序彻底崩溃
                
            sn, subtask_num = j1["sn"], j1.get("subtask_num", 1)

            for sub_no in range(1, subtask_num + 1):
                step2_data = [("tableBase", ""), ("isMerage", "1"), ("version", "v2"), ("sn", sn), ("sub_no", str(sub_no)), ("step", "2"), ("1", "1")]
                main_session.post(EXPORT_URL, headers=export_headers, data=step2_data, timeout=150)

            step3_data = [("tableBase", ""), ("isMerage", "1"), ("sn", sn), ("version", "v2"), ("step", "3"), ("1", "1")]
            r3 = main_session.post(EXPORT_URL, headers=export_headers, data=step3_data, timeout=150)
            task_id = r3.json().get("taskId")

            file_url = None
            # 增加轮询次数，确保大文件有足够时间生成 (30次 * 4秒 = 120秒轮询容忍度)
            for _ in range(30):
                step4_data = [("tableBase", ""), ("isMerage", "1"), ("sn", sn), ("version", "v2"), ("step", "4"), ("taskId", str(task_id)), ("1", "1")]
                j4 = main_session.post(EXPORT_URL, headers=export_headers, data=step4_data, timeout=150).json()
                if j4.get("success") and j4.get("file_url"):
                    file_url = j4["file_url"]
                    break
                time.sleep(4)
                
            if not file_url: 
                print(f"     ❌ [警告] 批次 {batch_idx + 1} 导出轮询超时！")
                continue

            # 核心修复 3：为保存的文件名增加批次前缀，防止同名覆盖
            original_filename = unquote(file_url.split("/")[-1])
            save_path = os.path.join(output_dir, f"batch_{batch_idx + 1}_{original_filename}")
            
            with open(save_path, "wb") as f:
                f.write(main_session.get(file_url, timeout=150).content)
                
            try:
                df = pd.read_csv(save_path, encoding='utf-8-sig') if save_path.endswith('.csv') else pd.read_excel(save_path, header=0)
                logger.info(f"批次 {batch_idx + 1} 文件列数={df.shape[1]}，行数={df.shape[0]}，列名={list(df.columns)}")

                # 导出文件为宽表格式：col0=组合sku编码, col1=关联sku个数,
                # 之后每 2 列一对 (关联sku编码N, 关联sku数量N)，最多 14 对
                parsed_count = 0
                if df.shape[1] >= 4:
                    for row_idx, row in df.iterrows():
                        parent_sku = str(row.iloc[0]).strip()
                        if not parent_sku or parent_sku == 'nan':
                            continue
                        try:
                            total_count = int(float(str(row.iloc[1]).strip()))
                        except (ValueError, TypeError):
                            total_count = 0
                        children = {}
                        for col_i in range(2, df.shape[1] - 1, 2):
                            child_sku = str(row.iloc[col_i]).strip()
                            if not child_sku or child_sku == 'nan':
                                break
                            try:
                                qty = int(float(str(row.iloc[col_i + 1]).strip()))
                            except (ValueError, TypeError):
                                qty = 1
                            children[child_sku] = children.get(child_sku, 0) + qty
                        if children:
                            # 存 (子件dict, 关联sku个数)，件数以 col1 为准
                            final_combo_dict[parent_sku] = (children, total_count)
                            parsed_count += 1
                else:
                    # 兼容旧版单列 "关联sku信息" 格式
                    combo_col = [c for c in df.columns if '组合sku编码' in str(c)]
                    child_col = [c for c in df.columns if '关联sku信息' in str(c)]
                    if combo_col and child_col:
                        for _, row in df.iterrows():
                            parent_sku = str(row[combo_col[0]]).strip()
                            child_info_str = str(row[child_col[0]]).strip()
                            if parent_sku and parent_sku != 'nan' and child_info_str and child_info_str != 'nan':
                                children = {}
                                for item in re.split(r'[,;\n\r]+', child_info_str):
                                    item = item.strip()
                                    if not item: continue
                                    match = re.search(r'^(.+?)[\*xX]\s*(\d+)$', item)
                                    if match: child_sku, qty = match.group(1).strip(), int(match.group(2))
                                    else: child_sku, qty = item.strip(), 1
                                    children[child_sku] = children.get(child_sku, 0) + qty
                                if children:
                                    total_count = sum(children.values())
                                    final_combo_dict[parent_sku] = (children, total_count)
                                    parsed_count += 1
                    else:
                        logger.warning(f"批次 {batch_idx + 1} 未找到有效列：combo_col={combo_col}, child_col={child_col}")

                logger.info(f"批次 {batch_idx + 1} 解析完成，本批提取组合结构 {parsed_count} 条")
            except Exception as e:
                logger.error(f"Pandas 解析批次 {batch_idx + 1} 组合文件失败: {e}", exc_info=True)
                
        except requests.exceptions.RequestException as re_err:
            print(f"     ❌ [网络超时] 批次 {batch_idx + 1} 请求过于庞大导致无响应: {re_err}")
        except Exception as e:
            print(f"     ❌ [未知崩溃] 批次 {batch_idx + 1} 导出解析阶段发生异常: {e}")

    return final_combo_dict

def _resolve_price(item, sku_label):
    try: purchase_price = float(item.get("purchasePrice", 0.0) or 0.0)
    except: purchase_price = 0.0
    try: default_cost = float(item.get("defaultCost", 0.0) or 0.0)
    except: default_cost = 0.0
    if purchase_price > 0: return purchase_price, ""
    if default_cost > 0: return default_cost, "兜底"
    return 0.0, "均为0"

async def batch_fetch_stock_prices(session, sku_batch):
    sku_data = "\r\n".join([s.strip() for s in sku_batch if s.strip()])

    result = {s: None for s in sku_batch}
    upper_index = {s.strip().upper(): s for s in sku_batch}

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            # FormData 必须在每次请求前重建，否则流已消费无法重用
            form_data = aiohttp.FormData()
            form_data.add_field("searchKey", "Stock_stockSku")
            form_data.add_field("operate", "=")
            form_data.add_field("isBatchSearch", "1")
            form_data.add_field("selecttype", "stockSku")
            form_data.add_field("stockData", sku_data)
            form_data.add_field("page", "1")
            form_data.add_field("rowsPerPage", "200")
            async with session.post(STOCK_API_URL, data=form_data, timeout=TIMEOUT_SETTING) as resp:
                if resp.status != 200:
                    logger.warning(f"库存取价批次 HTTP {resp.status}，第{attempt}次重试")
                    continue
                text = await resp.text()
                res_json = json.loads(text[text.find('{'):text.rfind('}') + 1])
                if not res_json.get("success") or not res_json.get("hasData"):
                    return result
                stock_items = res_json.get("stockData", [])
                for item in stock_items:
                    sku_key = str(item.get("stockSku", "")).strip().upper()
                    if sku_key in upper_index:
                        orig_sku = upper_index[sku_key]
                        if result[orig_sku] is None:
                            result[orig_sku] = _resolve_price(item, orig_sku)
                logger.info(f"库存取价批次: 发送{len(sku_batch)}条，API返回{len(stock_items)}行，命中{sum(1 for v in result.values() if v is not None)}条")
                return result
        except Exception as exc:
            logger.warning(f"库存取价批次异常(第{attempt}次): {exc}", exc_info=(attempt == MAX_RETRIES))
            if attempt < MAX_RETRIES: await asyncio.sleep(2 * attempt)
    return result


# ==========================================
# 💾 4. 双流缓存写入器
# ==========================================
class CacheWriter:
    def __init__(self, data_path, error_path):
        self.data_path = data_path
        self.error_path = error_path
        self.data_buffer = []
        self.error_buffer = []

    async def add(self, msku, result):
        async with file_lock:
            if isinstance(result, (int, float)): self.data_buffer.append((msku, result))
            else: self.error_buffer.append((msku, result))
            if max(len(self.data_buffer), len(self.error_buffer)) >= WRITE_FLUSH_EVERY:
                self._flush_locked()

    def _flush_locked(self):
        if self.data_buffer:
            with open(self.data_path, 'a', encoding='utf-8-sig', newline='') as f:
                for msku, cost in self.data_buffer: f.write(f"{msku},{cost}\n")
            self.data_buffer.clear()
        if self.error_buffer:
            with open(self.error_path, 'a', encoding='utf-8-sig', newline='') as f:
                for msku, reason in self.error_buffer: f.write(f"{msku},{reason}\n")
            self.error_buffer.clear()

    async def flush(self):
        async with file_lock: self._flush_locked()


# ==========================================
# 🚀 5. 主流水线集成引擎
# ==========================================
async def run_pipeline(task_list, main_session, amz_cookie_str, memcache_key, cache_file, error_cache_file, output_dir):
    amz_headers = {
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "User-Agent": "Mozilla/5.0",
        "X-Requested-With": "XMLHttpRequest",
        "Origin": AMZ_BASE_URL, 
        "Cookie": amz_cookie_str
    }
    connector = aiohttp.TCPConnector(limit=CONCURRENCY_LIMIT)
    semaphore = asyncio.Semaphore(CONCURRENCY_LIMIT)
    writer = CacheWriter(cache_file, error_cache_file)

    cleaned = [(m.strip(), l.strip()) for m, l in task_list]
    msku_to_local = {m: l for m, l in cleaned}
    all_mskus = [m for m, _ in cleaned]
    final_results = {}
    
    try:
        async with aiohttp.ClientSession(headers=amz_headers, connector=connector) as amz_sess:

            # ---------- 阶段 A ----------
            print(f"  🔎 阶段 A: 批量判别 SKU 类型 (每批 {BATCH_SIZE_TYPE} 条)...")
            type_map = {}
            async def _do_type_batch(batch):
                async with semaphore: return await batch_check_sku_types(amz_sess, batch)
                
            batches = [all_mskus[i:i + BATCH_SIZE_TYPE] for i in range(0, len(all_mskus), BATCH_SIZE_TYPE)]
            for fut in tqdm(asyncio.as_completed([_do_type_batch(b) for b in batches]), total=len(batches), desc="A.类型判定"):
                type_map.update(await fut)

            normal_mskus = [m for m, t in type_map.items() if t == "普通SKU"]
            combo_mskus = [m for m, t in type_map.items() if t == "组合SKU"]
            for m, t in type_map.items():
                if t in ["无此SKU", "判断超时"]: final_results[m] = "库存SKU获取不到采购价"

            # 导出组合明细表
            if combo_mskus:
                combo_export_data = [{"MSKU": m, "组合SKU": msku_to_local[m]} for m in combo_mskus]
                combo_export_path = os.path.join(output_dir, "阶段A_组合SKU清单.xlsx")
                try:
                    df_combo = pd.DataFrame(combo_export_data).drop_duplicates(subset=["MSKU", "组合SKU"])
                    df_combo.to_excel(combo_export_path, index=False)
                    print(f"     ↳ [导出] 组合SKU明细已导出，共 {len(df_combo)} 行双列数据，保存至: {combo_export_path}")
                except Exception as e:
                    logger.error(f"导出阶段A组合SKU清单失败: {e}")

            # ---------- 阶段 C1 ----------
            combo_local_skus_unique = list(set([msku_to_local[m] for m in combo_mskus]))
            combo_local_to_children = {}
            
            if combo_local_skus_unique:
                print(f"  🧩 阶段 C1: 发起 {len(combo_local_skus_unique)} 个独立组合SKU的导出解析任务...")
                combo_local_to_children = await asyncio.to_thread(
                    sync_export_combo_and_parse, 
                    main_session, 
                    memcache_key, 
                    combo_local_skus_unique, 
                    output_dir
                )
                print(f"     ↳ 成功提取底层子件结构 {len(combo_local_to_children)} 个")

            # ---------- 阶段 B ----------
            print(f"  💰 阶段 B: 批量获取底层库存采购价 (每批 {BATCH_SIZE_STOCK} 条)...")
            normal_local_skus = {msku_to_local[m] for m in normal_mskus}
            all_child_skus = set()
            for entry in combo_local_to_children.values():
                all_child_skus.update(entry[0].keys())

            all_query_skus = sorted(normal_local_skus | all_child_skus)
            price_cache = {}

            if all_query_skus:
                async def _do_stock_batch(batch):
                    async with semaphore: return await batch_fetch_stock_prices(amz_sess, batch)

                stock_batches = [all_query_skus[i:i + BATCH_SIZE_STOCK] for i in range(0, len(all_query_skus), BATCH_SIZE_STOCK)]
                for fut in tqdm(asyncio.as_completed([_do_stock_batch(b) for b in stock_batches]), total=len(stock_batches), desc="B.批量库存取价"):
                    price_cache.update(await fut)

                # 二次补查：对主批次仍为 None 的 SKU 缩小批次重试（20条/批）
                retry_skus = [s for s in all_query_skus if price_cache.get(s) is None]
                if retry_skus:
                    RETRY_BATCH = 20
                    print(f"     ↳ 二次补查 {len(retry_skus)} 个未命中SKU (每批 {RETRY_BATCH} 条)...")
                    retry_batches = [retry_skus[i:i + RETRY_BATCH] for i in range(0, len(retry_skus), RETRY_BATCH)]
                    for fut in tqdm(asyncio.as_completed([_do_stock_batch(b) for b in retry_batches]), total=len(retry_batches), desc="B.补查"):
                        for sku, val in (await fut).items():
                            if val is not None and price_cache.get(sku) is None:
                                price_cache[sku] = val
                    recovered = sum(1 for s in retry_skus if price_cache.get(s) is not None)
                    print(f"     ↳ 补查回收 {recovered}/{len(retry_skus)} 条")

            # ---------- 阶段 C2 ----------
            logger.info("========== 开始执行最后结算与异常排查 ==========")

            normal_success = normal_fail = 0
            for m in normal_mskus:
                local_sku = msku_to_local[m]
                p_data = price_cache.get(local_sku)
                if p_data and p_data[0] > 0.0:
                    final_results[m] = round(p_data[0] + 1.0, 2)
                    normal_success += 1
                else:
                    reason = f"采购价={p_data[0]}" if p_data is not None else "库存系统未找到(可能库存为0)"
                    logger.warning(f"普通件取价失败: MSKU={m}, 本地SKU={local_sku}, {reason}")
                    final_results[m] = "库存SKU获取不到采购价"
                    normal_fail += 1

            missing_local_skus = set(combo_local_skus_unique) - set(combo_local_to_children.keys())
            if missing_local_skus:
                logger.warning(f"C1 导出解析后，以下 {len(missing_local_skus)} 个本地SKU未提取到子件结构: {sorted(missing_local_skus)}")

            combo_success = combo_fail = 0
            for m in combo_mskus:
                ls = msku_to_local[m]
                entry = combo_local_to_children.get(ls)

                if not entry:
                    logger.warning(f"组合解析失败: MSKU={m}, 本地SKU={ls} 未在C1导出结果中找到子件 (C1共提取{len(combo_local_to_children)}条，缺失{len(missing_local_skus)}条)")
                    final_results[m] = "组合解析失败(未提取到子件)"
                    combo_fail += 1
                    continue

                children_dict, total_items_count = entry
                unit_price_sum = 0.0
                missing_child_skus = []

                for child_sku in children_dict:
                    p_data = price_cache.get(child_sku)
                    if p_data is None:
                        missing_child_skus.append(child_sku)
                    else:
                        unit_price_sum += p_data[0]

                if missing_child_skus:
                    logger.warning(f"MSKU={m} 本地SKU={ls} 子件在库存系统中未找到: {missing_child_skus}")
                    final_results[m] = "库存SKU获取不到采购价"
                    combo_fail += 1
                else:
                    # 公式：各子件单价之和 + 子件种类数
                    child_type_count = len(children_dict)
                    final_combo_price = unit_price_sum + child_type_count
                    final_results[m] = round(final_combo_price, 2)
                    combo_success += 1

            print(f"  📊 结算汇总:")
            print(f"     ↳ 普通件：成功 {normal_success} | 失败 {normal_fail}")
            print(f"     ↳ 组合件：成功 {combo_success} | 失败 {combo_fail}")

    finally:
        logger.info("======= 正在执行收尾刷盘操作 =======")
        for m, _ in cleaned: 
            if m in final_results:
                await writer.add(m, final_results[m])
        await writer.flush()

# ==========================================
# ✂️ 6. 拆分逻辑
# ==========================================
def auto_split_to_excel(csv_path, chunk_size=9900):
    if not os.path.exists(csv_path): return
    output_dir = os.path.dirname(csv_path)
    template_path = os.path.join(os.path.dirname(__file__), "成本导入模板.xlsx")

    print(f"\n  ✂️ 开始拆分为直接可上传的 Excel (每份 {chunk_size} 条)...")
    if not os.path.exists(template_path):
         print("  ❌ 找不到模板文件，跳过拆分。")
         return

    try:
        try: full_df = pd.read_csv(csv_path, encoding='utf-8-sig')
        except: full_df = pd.read_csv(csv_path, encoding='gbk')
        
        if full_df.empty:
            print("  ℹ️ 报表内无有效数据，无需拆分。")
            return
            
        data_records = full_df.iloc[:, :2].values.tolist()
        chunks = [data_records[i:i + chunk_size] for i in range(0, len(data_records), chunk_size)]
        
        for i, chunk in enumerate(chunks):
            target_file = os.path.join(output_dir, f"上传文件_第{i+1}部分.xlsx")
            wb = openpyxl.load_workbook(template_path)
            ws = wb.active
            start_row = 2
            for row_idx, row_data in enumerate(chunk):
                ws.cell(row=start_row + row_idx, column=1, value=row_data[0])
                ws.cell(row=start_row + row_idx, column=2, value=row_data[1])
            wb.save(target_file)
            wb.close()
            print(f"    -> 已生成分片: {os.path.basename(target_file)}")
    except Exception as e:
        print(f"  ❌ Excel 拆分失败: {e}")

# ==========================================
# 🏁 7. 入口函数 
# ==========================================
def run_cost(input_csv: str, output_dir: str):
    log_file = os.path.join(output_dir, "spider_debug.log")
    if logger.hasHandlers(): logger.handlers.clear()
    file_handler = FlushFileHandler(log_file, encoding="utf-8")
    file_handler.setFormatter(logging.Formatter('%(asctime)s [%(levelname)s] %(message)s'))
    logger.addHandler(file_handler)

    output_csv = os.path.join(output_dir, "最终采购成本表.csv")
    error_csv = os.path.join(output_dir, "不更新MSKU表.csv")        
    cache_file = os.path.join(output_dir, "task_cache.csv")
    error_cache_file = os.path.join(output_dir, "error_cache.csv")  

    # 【核心修复1】每次运行强制覆盖重写双流缓存文件，放弃断点续传逻辑，保证全量重跑
    with open(cache_file, 'w', encoding='utf-8-sig', newline='') as f: 
        f.write("MSKU,最终采购成本\n")
    with open(error_cache_file, 'w', encoding='utf-8-sig', newline='') as f: 
        f.write("MSKU,原因\n")

    print(f"  📦 加载输入表 {input_csv}...")
    try: df = pd.read_csv(input_csv, encoding='utf-8')
    except: df = pd.read_csv(input_csv, encoding='gb18030')

    df.columns = df.columns.str.strip()
    col_m, col_s, col_l = df.columns[0], df.columns[1], df.columns[2]
    
    df[col_m] = df[col_m].astype(str).str.strip()
    df[col_s] = df[col_s].astype(str).str.strip()
    df[col_l] = df[col_l].astype(str).str.strip()
    df = df.replace({'nan': None, '': None})

    out_of_stock_df = df[df[col_s] == '断货'].copy()
    valid_df = df[(df[col_s] != '断货')].dropna(subset=[col_m, col_l]).copy()
    
    # 彻底清理掉 old_cache 判断机制，直接转换为待处理任务
    pending_tasks = list(zip(valid_df[col_m], valid_df[col_l]))

    if not out_of_stock_df.empty:
        with open(error_cache_file, 'a', encoding='utf-8-sig', newline='') as f:
            for m in out_of_stock_df[col_m].tolist():
                f.write(f"{m},断货\n")

    print(f"  📊 待处理在售 SKU: {len(pending_tasks)} 条\n")

    if pending_tasks:
        main_session, amz_cookie_str, memcache_key = get_global_auth()
        if main_session and amz_cookie_str:
            try: 
                asyncio.run(run_pipeline(pending_tasks, main_session, amz_cookie_str, memcache_key, cache_file, error_cache_file, output_dir))
            except KeyboardInterrupt: 
                print("\n  ⚠️ 任务手动中断，进度已安全保存！")
            except Exception as e: 
                # 【核心修复2】如果内部再发生阻断性崩溃，将其显式打印在控制台
                print(f"\n  ❌ 流水线运行发生严重崩溃，任务已终止: {e} (详情见 spider_debug.log)")
                logger.error(f"全局异常: {e}", exc_info=True)

    print("\n  📝 生成最终报表...")
    if os.path.exists(cache_file): pd.read_csv(cache_file, header=0).to_csv(output_csv, index=False, encoding='utf-8-sig')
    if os.path.exists(error_cache_file): pd.read_csv(error_cache_file, header=0).to_csv(error_csv, index=False, encoding='utf-8-sig')
    print("  🎉 成本计算完毕！")
    
    auto_split_to_excel(output_csv, chunk_size=5100)